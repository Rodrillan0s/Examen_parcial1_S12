import unittest
import io
import openpyxl
from fastapi.testclient import TestClient

from app import create_app
from app.utils.security import create_access_token
from app.services.compras_lotes_services import (
    generar_plantilla_excel_lotes,
    analizar_y_previsualizar_excel,
    confirmar_e_importar_lote_db,
    listar_ordenes_compra_db,
    aprobar_orden_compra_db,
    recibir_mercaderia_orden_db
)

class TestComprasLotesImportacion(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        cls.client = TestClient(cls.app)
        cls.token = create_access_token(
            nro_usuario=1,
            username='admin',
            id_rol=1,
            id_empresa=1,
            nombre='Admin',
            apellido='Aurora',
            roles=['ADMINISTRADOR']
        )
        cls.headers = {'Authorization': f'Bearer {cls.token}'}
        cls.token_data = {
            'id_empresa': 1,
            'rol': 'ADMINISTRADOR',
            'nombre': 'Admin',
            'apellido': 'Aurora',
            'nro_usuario': 1
        }

    # --------------------------------------------------------------------------
    # 1. GENERACIÓN DE LA PLANTILLA OFICIAL EXCEL
    # --------------------------------------------------------------------------
    def test_01_generar_plantilla_excel(self):
        buf = generar_plantilla_excel_lotes(id_empresa=1)
        self.assertIsInstance(buf, io.BytesIO)
        content = buf.getvalue()
        self.assertGreater(len(content), 1000)

        wb = openpyxl.load_workbook(io.BytesIO(content))
        self.assertIn("Carga_Prendas", wb.sheetnames)
        self.assertIn("Catalogos_Referencia", wb.sheetnames)

        ws = wb["Carga_Prendas"]
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Código Producto *", headers)
        self.assertIn("Cantidad Lote *", headers)
        self.assertIn("Precio Venta (Bs.) *", headers)

    # --------------------------------------------------------------------------
    # 2. PRE-VISUALIZACIÓN Y VALIDACIÓN EN SECO (DRY-RUN)
    # --------------------------------------------------------------------------
    def test_02_previsualizar_excel_valido_e_invalido(self):
        # Crear un archivo Excel en memoria para simular la subida del usuario
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Carga_Prendas"
        ws.append([
            "codigo_producto", "nombre_producto", "categoria", "genero",
            "talla", "color", "sku", "cantidad_lote", "costo_unitario",
            "precio_venta", "stock_minimo", "proveedor", "numero_lote"
        ])
        # Fila 1: Válida
        ws.append(["TEST-VES-001", "Vestido Test Seda", "Vestidos", "Damas", "M", "Negro", "TEST-VES-M-NEG", 30, 80.0, 180.0, 5, "Proveedor Test", "LOT-TEST-01"])
        # Fila 2: Inválida (sin nombre y cantidad negativa)
        ws.append(["TEST-BLU-002", "", "Blusas", "Damas", "S", "Blanco", "TEST-BLU-S-BLA", -5, 50.0, 0.0, 5, "Proveedor Test", "LOT-TEST-01"])

        buf = io.BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()

        preview = analizar_y_previsualizar_excel(excel_bytes, id_empresa=1)
        self.assertEqual(preview["total_filas"], 2)
        self.assertEqual(preview["filas_validas"], 1)
        self.assertEqual(preview["filas_con_error"], 1)
        self.assertEqual(preview["total_prendas"], 30)
        self.assertEqual(preview["filas"][1]["estado"], "ERROR")

    # --------------------------------------------------------------------------
    # 3. CONFIRMACIÓN E IMPORTACIÓN A BASE DE DATOS E INVENTARIO
    # --------------------------------------------------------------------------
    def test_03_confirmar_importacion_afecta_inventario(self):
        cod_unico = "AUR-IMP-TEST-99"
        filas = [{
            "codigo_producto": cod_unico,
            "nombre_producto": "Vestido Coctel Importado",
            "categoria": "Vestidos",
            "genero": "Damas",
            "talla": "L",
            "color": "Verde Esmeralda",
            "sku": f"{cod_unico}-L-VER",
            "cantidad": 45,
            "costo_unitario": 90.0,
            "precio_venta": 220.0,
            "subtotal_costo": 4050.0,
            "stock_minimo": 5,
            "proveedor": "Importadora Textil Andina",
            "numero_lote": "LOT-UNITTEST-001"
        }]

        payload = {
            "id_sucursal": 1,
            "filas": filas,
            "generar_orden_compra": True,
            "numero_lote": "LOT-UNITTEST-001",
            "observaciones": "Test unitario de importación masiva"
        }

        res = confirmar_e_importar_lote_db(payload, self.token_data)
        self.assertTrue(res["success"])
        self.assertEqual(res["total_prendas"], 45)
        self.assertIsNotNone(res["id_lote"])
        self.assertIsNotNone(res["id_orden_compra"])

    # --------------------------------------------------------------------------
    # 4. CICLO DE VIDA DE ÓRDENES DE COMPRA (APROBACIÓN Y RECEPCIÓN)
    # --------------------------------------------------------------------------
    def test_04_flujo_orden_compra_y_recepcion(self):
        # 1. Crear Orden de Compra Pendiente
        ordenes = listar_ordenes_compra_db(id_empresa=1)
        self.assertIsInstance(ordenes, list)

        # Usar la primera orden o crear una
        from app.classes.postgres import PostgreSQL
        db = PostgreSQL()
        db.create_connection()
        try:
            q = """
            INSERT INTO comercio.t_orden_compra (
                id_empresa, id_sucursal, numero_orden, estado, total_estimado, id_usuario_creador
            ) VALUES (1, 1, 'OC-TEST-0099', 'PENDIENTE_APROBACION', 1500.0, (SELECT id_usuario FROM comercio.t_usuario WHERE id_empresa = 1 LIMIT 1))
            RETURNING id_orden_compra;
            """
            oc = db.execute_query(q, fetchone=True, commit=True)
            id_oc = oc[0]

            # Agregar un ítem
            q_it = """
            INSERT INTO comercio.t_detalle_orden_compra (
                id_orden_compra, codigo_producto, nombre_producto, talla, color,
                cantidad_solicitada, costo_unitario, subtotal
            ) VALUES (%s, 'TEST-PROD-OC', 'Prenda OC Test', 'M', 'Azul', 20, 75.0, 1500.0);
            """
            db.execute_query(q_it, (id_oc,), commit=True)
        finally:
            db.close_connection()

        # 2. Aprobar la Orden
        res_aprob = aprobar_orden_compra_db(id_oc, self.token_data)
        self.assertTrue(res_aprob["success"])

        # 3. Recibir la Mercadería en Sucursal
        res_rec = recibir_mercaderia_orden_db(id_oc, {"observaciones": "Recepción conforme"}, self.token_data)
        self.assertTrue(res_rec["success"])
        self.assertEqual(res_rec["total_prendas"], 20)

    # --------------------------------------------------------------------------
    # 5. ENDPOINTS REST FASTAPI
    # --------------------------------------------------------------------------
    def test_05_endpoint_descarga_plantilla(self):
        resp = self.client.get("/api/inventario/importacion/plantilla", headers=self.headers)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument", resp.headers["content-type"])

    def test_06_endpoint_listar_ordenes_y_lotes(self):
        r_oc = self.client.get("/api/compras/ordenes", headers=self.headers)
        self.assertEqual(r_oc.status_code, 200)
        self.assertTrue(r_oc.json()["success"])

        r_lot = self.client.get("/api/compras/lotes", headers=self.headers)
        self.assertEqual(r_lot.status_code, 200)
        self.assertTrue(r_lot.json()["success"])


if __name__ == '__main__':
    unittest.main()

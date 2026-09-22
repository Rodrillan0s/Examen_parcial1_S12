import io
import datetime
import random
from decimal import Decimal
from typing import Dict, Any, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.classes.postgres import PostgreSQL
from app.config import Config
from app.repos.bitacora_repos import registrar_evento_db

def _get_schema() -> str:
    return Config.SCHEMA or 'comercio'

# ==============================================================================
# 1. GENERADOR DE PLANTILLA FIJA OFICIAL EXCEL (.XLSX)
# ==============================================================================

def generar_plantilla_excel_lotes(id_empresa: int) -> io.BytesIO:
    """
    Genera un archivo Excel (.xlsx) estructurado con dos hojas:
    - Hoja 1: 'Carga_Prendas' con las columnas fijas y ejemplos ilustrativos.
    - Hoja 2: 'Catalogos_Referencia' con categorías, tallas y colores actuales.
    """
    wb = openpyxl.Workbook()
    schema = _get_schema()

    # --- HOJA 1: Carga_Prendas ---
    ws = wb.active
    ws.title = "Carga_Prendas"

    # Estilos
    header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    example_font = Font(name="Calibri", size=10, italic=True, color="475569")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    headers = [
        ("codigo_producto", "Código Producto *", 18, align_left),
        ("nombre_producto", "Nombre / Modelo Prenda *", 30, align_left),
        ("categoria", "Categoría *", 18, align_left),
        ("genero", "Género (Damas/Caballeros/Unisex)", 22, align_center),
        ("talla", "Talla (S/M/L/XL/etc) *", 16, align_center),
        ("color", "Color *", 16, align_left),
        ("sku", "SKU (Opcional)", 20, align_left),
        ("cantidad_lote", "Cantidad Lote *", 16, align_right),
        ("costo_unitario", "Costo Unitario (Bs.) *", 20, align_right),
        ("precio_venta", "Precio Venta (Bs.) *", 18, align_right),
        ("stock_minimo", "Stock Mínimo Alerta", 18, align_right),
        ("proveedor", "Proveedor (Razón Social / NIT)", 28, align_left),
        ("numero_lote", "N° Lote (Opcional)", 20, align_left)
    ]

    # Escribir encabezados
    for col_num, (_, title, width, alignment) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = alignment
        ws.column_dimensions[get_column_letter(col_num)].width = width

    # Filas de ejemplo real
    ejemplos = [
        ("AUR-BLU-001", "Blusa Seda Italiana Escote V", "Blusas", "Damas", "S", "Blanco Marfil", "BLU-SED-S-BLA", 40, 75.00, 160.00, 5, "Textiles del Valle S.R.L.", "LOT-2026-001"),
        ("AUR-BLU-001", "Blusa Seda Italiana Escote V", "Blusas", "Damas", "M", "Negro Ónix", "BLU-SED-M-NEG", 60, 75.00, 160.00, 5, "Textiles del Valle S.R.L.", "LOT-2026-001"),
        ("AUR-VES-002", "Vestido Noche Gala Terciopelo", "Vestidos", "Damas", "M", "Rojo Rubí", "VES-GAL-M-ROJ", 25, 140.00, 320.00, 4, "Confecciones Alta Costura", "LOT-2026-001"),
        ("AUR-PAN-003", "Pantalón Denim Flare Cintura Alta", "Pantalones", "Damas", "38", "Azul Marino", "PAN-DEN-38-AZU", 50, 95.00, 210.00, 8, "Importadora Denim Jeans", "LOT-2026-001"),
    ]

    for row_idx, data_row in enumerate(ejemplos, start=2):
        for col_idx, val in enumerate(data_row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = example_font
            cell.border = thin_border
            # Alineación según columna
            cell.alignment = headers[col_idx - 1][3]

    # --- HOJA 2: Catalogos_Referencia ---
    ws2 = wb.create_sheet(title="Catalogos_Referencia")
    ws2.cell(row=1, column=1, value="Categorías Existentes").font = Font(bold=True)
    ws2.cell(row=1, column=3, value="Tallas Registradas").font = Font(bold=True)
    ws2.cell(row=1, column=5, value="Colores Registrados").font = Font(bold=True)

    db = PostgreSQL()
    db.create_connection()
    try:
        # Categorías de la empresa
        q_cat = f"SELECT nombre FROM {schema}.t_categoria WHERE (id_empresa = %s OR id_empresa IS NULL) AND activo = TRUE ORDER BY nombre;"
        cats = db.execute_query(q_cat, (id_empresa,), fetchall=True) or []
        for i, c in enumerate(cats, start=2):
            ws2.cell(row=i, column=1, value=c[0])

        # Tallas
        q_talla = f"SELECT nombre FROM {schema}.t_talla WHERE activo = TRUE ORDER BY id_talla;"
        tallas = db.execute_query(q_talla, fetchall=True) or []
        for i, t in enumerate(tallas, start=2):
            ws2.cell(row=i, column=3, value=t[0])

        # Colores
        q_color = f"SELECT nombre FROM {schema}.t_color WHERE activo = TRUE ORDER BY nombre;"
        colores = db.execute_query(q_color, fetchall=True) or []
        for i, cl in enumerate(colores, start=2):
            ws2.cell(row=i, column=5, value=cl[0])

        ws2.column_dimensions["A"].width = 25
        ws2.column_dimensions["C"].width = 25
        ws2.column_dimensions["E"].width = 25
    finally:
        db.close_connection()

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ==============================================================================
# 2. ANALIZADOR Y PRE-VISUALIZACIÓN (PREVIEW / DRY-RUN)
# ==============================================================================

def analizar_y_previsualizar_excel(file_content: bytes, id_empresa: int, id_sucursal: Optional[int] = None) -> Dict[str, Any]:
    """
    Analiza un archivo Excel cargado en memoria, valida cada fila y retorna
    un reporte diagnóstico exhaustivo (Dry-Run) con la clasificación de cada registro.
    """
    schema = _get_schema()
    db = PostgreSQL()
    db.create_connection()
    try:
        # Cargar catálogos en memoria para validación O(1)
        q_prods = f"SELECT codigo_producto, id_producto, nombre FROM {schema}.t_producto WHERE id_empresa = %s;"
        prods_raw = db.execute_query(q_prods, (id_empresa,), fetchall=True) or []
        map_prods = {str(p[0]).strip().upper(): {"id": p[1], "nombre": p[2]} for p in prods_raw if p[0]}

        q_cats = f"SELECT id_categoria, LOWER(nombre) FROM {schema}.t_categoria WHERE (id_empresa = %s OR id_empresa IS NULL);"
        cats_raw = db.execute_query(q_cats, (id_empresa,), fetchall=True) or []
        map_cats = {c[1].strip(): c[0] for c in cats_raw if c[1]}

        q_tallas = f"SELECT id_talla, LOWER(nombre) FROM {schema}.t_talla;"
        tallas_raw = db.execute_query(q_tallas, fetchall=True) or []
        map_tallas = {t[1].strip(): t[0] for t in tallas_raw if t[1]}

        q_colores = f"SELECT id_color, LOWER(nombre) FROM {schema}.t_color;"
        colores_raw = db.execute_query(q_colores, fetchall=True) or []
        map_colores = {cl[1].strip(): cl[0] for cl in colores_raw if cl[1]}

        # Parsear Excel
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_content), data_only=True)
        sheet = wb.active

        filas_resultado = []
        total_prendas = 0
        costo_total = Decimal("0.0")
        valor_venta_total = Decimal("0.0")
        prods_nuevos_count = 0
        prods_existentes_count = 0
        filas_con_error = 0

        # Leer encabezados (fila 1)
        headers = [str(cell.value or '').strip().lower() for cell in sheet[1]]
        
        # Mapeo de índices de columnas
        def find_col_idx(aliases: List[str]) -> Optional[int]:
            for alias in aliases:
                for idx, h in enumerate(headers):
                    if alias in h:
                        return idx
            return None

        idx_cod = find_col_idx(["codigo_producto", "código producto", "codigo"])
        idx_nom = find_col_idx(["nombre_producto", "nombre", "modelo"])
        idx_cat = find_col_idx(["categoria", "categoría"])
        idx_gen = find_col_idx(["genero", "género"])
        idx_tal = find_col_idx(["talla"])
        idx_col = find_col_idx(["color"])
        idx_sku = find_col_idx(["sku"])
        idx_can = find_col_idx(["cantidad_lote", "cantidad"])
        idx_cos = find_col_idx(["costo_unitario", "costo"])
        idx_pre = find_col_idx(["precio_venta", "precio"])
        idx_min = find_col_idx(["stock_minimo", "minimo", "mínimo"])
        idx_pro = find_col_idx(["proveedor"])
        idx_lot = find_col_idx(["numero_lote", "lote"])

        if idx_cod is None or idx_nom is None or idx_can is None:
            raise ValueError("El archivo Excel no cuenta con las columnas mínimas requeridas ('codigo_producto', 'nombre_producto', 'cantidad_lote').")

        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            # Omitir filas completamente vacías
            if not any(row):
                continue

            cod = str(row[idx_cod] or '').strip().upper() if idx_cod is not None and idx_cod < len(row) else ''
            nom = str(row[idx_nom] or '').strip() if idx_nom is not None and idx_nom < len(row) else ''
            cat = str(row[idx_cat] or '').strip() if idx_cat is not None and idx_cat < len(row) else 'General'
            gen = str(row[idx_gen] or '').strip() if idx_gen is not None and idx_gen < len(row) else 'Unisex'
            tal = str(row[idx_tal] or '').strip() if idx_tal is not None and idx_tal < len(row) else 'Única'
            col = str(row[idx_col] or '').strip() if idx_col is not None and idx_col < len(row) else 'Estándar'
            sku = str(row[idx_sku] or '').strip() if idx_sku is not None and idx_sku < len(row) else ''
            
            raw_cant = row[idx_can] if idx_can is not None and idx_can < len(row) else 0
            raw_costo = row[idx_cos] if idx_cos is not None and idx_cos < len(row) else 0.0
            raw_precio = row[idx_pre] if idx_pre is not None and idx_pre < len(row) else 0.0
            raw_min = row[idx_min] if idx_min is not None and idx_min < len(row) else 5
            prov = str(row[idx_pro] or '').strip() if idx_pro is not None and idx_pro < len(row) else ''
            lote = str(row[idx_lot] or '').strip() if idx_lot is not None and idx_lot < len(row) else ''

            # Validación de datos
            errores = []
            if not cod:
                errores.append("Falta código de producto")
            if not nom:
                errores.append("Falta nombre del producto")

            try:
                cant = int(raw_cant or 0)
                if cant <= 0:
                    errores.append("Cantidad debe ser mayor a 0")
            except:
                cant = 0
                errores.append("Cantidad inválida")

            try:
                costo = float(Decimal(str(raw_costo or 0)))
                if costo < 0:
                    errores.append("Costo unitario no puede ser negativo")
            except:
                costo = 0.0
                errores.append("Costo unitario inválido")

            try:
                precio = float(Decimal(str(raw_precio or 0)))
                if precio <= 0:
                    errores.append("Precio de venta debe ser mayor a 0")
            except:
                precio = 0.0
                errores.append("Precio de venta inválido")

            try:
                minimo = int(raw_min or 5)
            except:
                minimo = 5

            if not sku and cod:
                sku = f"{cod}-{tal[:3].upper()}-{col[:3].upper()}".replace(" ", "")

            # Clasificación de estado
            es_nuevo = cod not in map_prods
            cat_existe = cat.lower() in map_cats
            tal_existe = tal.lower() in map_tallas
            col_existe = col.lower() in map_colores

            if errores:
                filas_con_error += 1
                estado_fila = "ERROR"
            elif es_nuevo:
                prods_nuevos_count += 1
                estado_fila = "NUEVO_PRODUCTO"
            else:
                prods_existentes_count += 1
                estado_fila = "PRODUCTO_EXISTENTE"

            if not errores:
                total_prendas += cant
                costo_total += Decimal(str(cant)) * Decimal(str(costo))
                valor_venta_total += Decimal(str(cant)) * Decimal(str(precio))

            filas_resultado.append({
                "fila_excel": row_idx,
                "codigo_producto": cod,
                "nombre_producto": nom,
                "categoria": cat,
                "categoria_nueva": not cat_existe,
                "genero": gen,
                "talla": tal,
                "talla_nueva": not tal_existe,
                "color": col,
                "color_nuevo": not col_existe,
                "sku": sku,
                "cantidad": cant,
                "costo_unitario": costo,
                "precio_venta": precio,
                "subtotal_costo": round(cant * costo, 2),
                "stock_minimo": minimo,
                "proveedor": prov,
                "numero_lote": lote,
                "estado": estado_fila,
                "errores": errores
            })

        return {
            "valido": filas_con_error == 0 and len(filas_resultado) > 0,
            "total_filas": len(filas_resultado),
            "filas_validas": len(filas_resultado) - filas_con_error,
            "filas_con_error": filas_con_error,
            "total_prendas": total_prendas,
            "costo_total": float(costo_total),
            "valor_venta_total": float(valor_venta_total),
            "margen_bruto_estimado": float(valor_venta_total - costo_total),
            "productos_nuevos": prods_nuevos_count,
            "productos_existentes": prods_existentes_count,
            "filas": filas_resultado
        }
    finally:
        db.close_connection()


# ==============================================================================
# 3. CONFIRMACIÓN E IMPORTACIÓN TRANSACCIONAL A LA BASE DE DATOS
# ==============================================================================

def confirmar_e_importar_lote_db(
    payload: Dict[str, Any],
    token_data: Dict[str, Any]
) -> Dict[str, Any]:
    """
    Inserta transaccionalmente en PostgreSQL:
    1. Si corresponde, crea o asocia Orden de Compra (t_orden_compra y t_detalle_orden_compra).
    2. Registra el lote en t_lote.
    3. Asegura categorías, tallas, colores, productos y variantes.
    4. Afecta t_inventario (stock_actual += cant, stock_disponible += cant) en la sucursal seleccionada.
    5. Registra auditoría en t_movimiento_inventario (tipo_movimiento='ENTRADA_COMPRA').
    """
    id_empresa = int(
        payload.get("id_empresa")
        or token_data.get("id_empresa")
        or token_data.get("empresa_id")
        or 1
    )
    id_usuario = int(token_data.get("nro_usuario") or token_data.get("id_usuario") or 1)

    id_sucursal = payload.get("id_sucursal")
    if not id_sucursal:
        raise ValueError("Debe especificar la sucursal de destino para el ingreso del inventario.")
    id_sucursal = int(id_sucursal)

    filas = payload.get("filas") or []
    if not filas:
        raise ValueError("No se enviaron filas válidas para importar.")

    generar_orden = bool(payload.get("generar_orden_compra", True))
    guia_remision = payload.get("guia_remision") or ""
    observaciones = payload.get("observaciones") or "Importación masiva mediante archivo Excel"

    # Generar N° de Lote unificado si no viene
    numero_lote = payload.get("numero_lote")
    if not numero_lote:
        timestamp_lote = datetime.datetime.now().strftime("%Y%m%d_%H%M")
        rand_suf = random.randint(100, 999)
        numero_lote = f"LOT-{timestamp_lote}-{rand_suf}"

    schema = _get_schema()
    db = PostgreSQL()
    db.create_connection()
    try:
        # Validar que la sucursal pertenezca a la empresa
        q_suc = f"SELECT id_sucursal, nombre FROM {schema}.t_sucursal WHERE id_sucursal = %s AND id_empresa = %s AND activo = TRUE;"
        suc = db.execute_query(q_suc, (id_sucursal, id_empresa), fetchone=True)
        if not suc:
            raise ValueError(f"La sucursal ID {id_sucursal} no existe o no pertenece a la empresa {id_empresa}.")
        nombre_sucursal = suc[1]

        # Resolver id_usuario seguro
        q_u = f"SELECT id_usuario FROM {schema}.t_usuario WHERE id_usuario = %s;"
        u_chk = db.execute_query(q_u, (id_usuario,), fetchone=True)
        if not u_chk:
            q_u_alt = f"SELECT id_usuario FROM {schema}.t_usuario WHERE id_empresa = %s LIMIT 1;"
            u_alt = db.execute_query(q_u_alt, (id_empresa,), fetchone=True)
            id_usuario = u_alt[0] if u_alt else None

        # 1. Resolver Proveedor si viene especificado
        id_proveedor = payload.get("id_proveedor")
        if not id_proveedor and filas[0].get("proveedor"):
            prov_nom = filas[0]["proveedor"].strip()
            q_prov_find = f"SELECT id_proveedor FROM {schema}.t_proveedor WHERE LOWER(razon_social) = LOWER(%s) LIMIT 1;"
            prov_row = db.execute_query(q_prov_find, (prov_nom,), fetchone=True)
            if prov_row:
                id_proveedor = prov_row[0]
            else:
                q_prov_ins = f"INSERT INTO {schema}.t_proveedor (razon_social, estado) VALUES (%s, TRUE) RETURNING id_proveedor;"
                ins_p = db.execute_query(q_prov_ins, (prov_nom,), fetchone=True, commit=False)
                id_proveedor = ins_p[0] if ins_p else None

        total_prendas_lote = sum(int(f.get("cantidad", 0)) for f in filas)
        costo_total_lote = sum(Decimal(str(f.get("subtotal_costo", 0.0))) for f in filas)

        # 2. Crear Orden de Compra si corresponde
        id_orden_compra = None
        if generar_orden:
            numero_oc = f"OC-{datetime.datetime.now().strftime('%Y%m')}-{random.randint(1000, 9999)}"
            q_oc = f"""
            INSERT INTO {schema}.t_orden_compra (
                id_empresa, id_sucursal, id_proveedor, numero_orden,
                estado, total_estimado, observaciones, id_usuario_creador,
                id_usuario_aprobador, fecha_aprobacion
            ) VALUES (%s, %s, %s, %s, 'RECIBIDA_TOTAL', %s, %s, %s, %s, CURRENT_TIMESTAMP)
            RETURNING id_orden_compra;
            """
            oc_res = db.execute_query(
                q_oc,
                (id_empresa, id_sucursal, id_proveedor, numero_oc, float(costo_total_lote), observaciones, id_usuario, id_usuario),
                fetchone=True, commit=False
            )
            id_orden_compra = oc_res[0] if oc_res else None

        # 3. Registrar Lote en t_lote
        q_lote = f"""
        INSERT INTO {schema}.t_lote (
            id_empresa, id_sucursal, id_orden_compra, id_proveedor,
            numero_lote, guia_remision, total_prendas, costo_total_lote,
            id_usuario_receptor, observaciones
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id_lote;
        """
        lote_res = db.execute_query(
            q_lote,
            (id_empresa, id_sucursal, id_orden_compra, id_proveedor, numero_lote, guia_remision, total_prendas_lote, float(costo_total_lote), id_usuario, observaciones),
            fetchone=True, commit=False
        )
        id_lote = lote_res[0] if lote_res else None

        # 4. Procesar cada fila de prenda
        items_procesados = 0
        productos_creados = 0
        variantes_creadas = 0

        for f in filas:
            cod_prod = str(f.get("codigo_producto", "")).strip().upper()
            nom_prod = str(f.get("nombre_producto", "")).strip()
            cat_nom = str(f.get("categoria", "")).strip() or "General"
            gen_val = str(f.get("genero", "Unisex")).strip()
            talla_nom = str(f.get("talla", "Única")).strip()
            color_nom = str(f.get("color", "Estándar")).strip()
            sku_val = str(f.get("sku", "")).strip()
            cantidad = int(f.get("cantidad", 0))
            costo_unit = float(f.get("costo_unitario", 0.0))
            precio_vta = float(f.get("precio_venta", 0.0))
            stock_min = int(f.get("stock_minimo", 5))

            if cantidad <= 0:
                continue

            # A. Resolver o crear Categoría
            q_cat_find = f"SELECT id_categoria FROM {schema}.t_categoria WHERE LOWER(nombre) = LOWER(%s) AND (id_empresa = %s OR id_empresa IS NULL) LIMIT 1;"
            c_row = db.execute_query(q_cat_find, (cat_nom, id_empresa), fetchone=True)
            if c_row:
                id_cat = c_row[0]
            else:
                q_cat_ins = f"INSERT INTO {schema}.t_categoria (nombre, id_empresa, activo) VALUES (%s, %s, TRUE) RETURNING id_categoria;"
                c_ins = db.execute_query(q_cat_ins, (cat_nom, id_empresa), fetchone=True, commit=False)
                id_cat = c_ins[0]

            # B. Resolver o crear Talla
            q_tal_find = f"SELECT id_talla FROM {schema}.t_talla WHERE LOWER(nombre) = LOWER(%s) LIMIT 1;"
            t_row = db.execute_query(q_tal_find, (talla_nom,), fetchone=True)
            if t_row:
                id_tal = t_row[0]
            else:
                q_tal_ins = f"INSERT INTO {schema}.t_talla (nombre, id_empresa, activo, estado) VALUES (%s, %s, TRUE, TRUE) RETURNING id_talla;"
                t_ins = db.execute_query(q_tal_ins, (talla_nom, id_empresa), fetchone=True, commit=False)
                id_tal = t_ins[0]

            # C. Resolver o crear Color
            q_col_find = f"SELECT id_color FROM {schema}.t_color WHERE LOWER(nombre) = LOWER(%s) LIMIT 1;"
            cl_row = db.execute_query(q_col_find, (color_nom,), fetchone=True)
            if cl_row:
                id_col = cl_row[0]
            else:
                q_col_ins = f"INSERT INTO {schema}.t_color (nombre, id_empresa, codigo_hex, activo, estado) VALUES (%s, %s, '#475569', TRUE, TRUE) RETURNING id_color;"
                cl_ins = db.execute_query(q_col_ins, (color_nom, id_empresa), fetchone=True, commit=False)
                id_col = cl_ins[0]

            # D. Resolver o crear Producto Padre
            q_p_find = f"SELECT id_producto, precio FROM {schema}.t_producto WHERE codigo_producto = %s AND id_empresa = %s LIMIT 1;"
            p_row = db.execute_query(q_p_find, (cod_prod, id_empresa), fetchone=True)
            if p_row:
                id_producto = p_row[0]
            else:
                q_p_ins = f"""
                INSERT INTO {schema}.t_producto (
                    codigo_producto, nombre, id_categoria, genero,
                    precio, id_empresa, activo, estado, fecha_registro
                ) VALUES (%s, %s, %s, %s, %s, %s, TRUE, TRUE, CURRENT_TIMESTAMP)
                RETURNING id_producto;
                """
                p_ins = db.execute_query(q_p_ins, (cod_prod, nom_prod, id_cat, gen_val, precio_vta, id_empresa), fetchone=True, commit=False)
                id_producto = p_ins[0]
                productos_creados += 1

            # E. Resolver o crear Variante (t_producto_talla_color)
            if not sku_val:
                sku_val = f"{cod_prod}-{talla_nom[:3].upper()}-{color_nom[:3].upper()}".replace(" ", "")

            q_var_find = f"SELECT id_variante FROM {schema}.t_producto_talla_color WHERE id_producto = %s AND id_talla = %s AND id_color = %s LIMIT 1;"
            var_row = db.execute_query(q_var_find, (id_producto, id_tal, id_col), fetchone=True)
            if var_row:
                id_variante = var_row[0]
            else:
                q_var_ins = f"""
                INSERT INTO {schema}.t_producto_talla_color (
                    id_producto, id_talla, id_color, sku, codigo_barras, precio, activo, estado
                ) VALUES (%s, %s, %s, %s, %s, %s, TRUE, TRUE)
                RETURNING id_variante;
                """
                var_ins = db.execute_query(q_var_ins, (id_producto, id_tal, id_col, sku_val, sku_val, precio_vta), fetchone=True, commit=False)
                id_variante = var_ins[0]
                variantes_creadas += 1

            # F. Detalle de Orden de Compra si se generó
            if id_orden_compra:
                q_det_oc = f"""
                INSERT INTO {schema}.t_detalle_orden_compra (
                    id_orden_compra, id_producto, id_variante, codigo_producto,
                    nombre_producto, talla, color, sku, cantidad_solicitada,
                    cantidad_recibida, costo_unitario, subtotal
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                """
                db.execute_query(
                    q_det_oc,
                    (id_orden_compra, id_producto, id_variante, cod_prod, nom_prod, talla_nom, color_nom, sku_val, cantidad, cantidad, costo_unit, round(cantidad * costo_unit, 2)),
                    commit=False
                )

            # G. Afectación a Inventario en la Sucursal Destino (t_inventario)
            q_inv_find = f"SELECT id_inventario, stock_actual, stock_disponible FROM {schema}.t_inventario WHERE id_sucursal = %s AND id_variante = %s LIMIT 1;"
            inv_row = db.execute_query(q_inv_find, (id_sucursal, id_variante), fetchone=True)

            if inv_row:
                id_inventario = inv_row[0]
                stock_ant = inv_row[1]
                stock_nuevo = stock_ant + cantidad
                q_inv_upd = f"""
                UPDATE {schema}.t_inventario
                SET stock_actual = stock_actual + %s,
                    stock_disponible = stock_disponible + %s,
                    stock_minimo = %s,
                    fecha_actualizacion = CURRENT_TIMESTAMP
                WHERE id_inventario = %s;
                """
                db.execute_query(q_inv_upd, (cantidad, cantidad, stock_min, id_inventario), commit=False)
            else:
                stock_ant = 0
                stock_nuevo = cantidad
                q_inv_ins = f"""
                INSERT INTO {schema}.t_inventario (
                    id_sucursal, id_variante, stock_actual, stock_reservado,
                    stock_disponible, stock_minimo, estado, fecha_actualizacion
                ) VALUES (%s, %s, %s, 0, %s, %s, TRUE, CURRENT_TIMESTAMP)
                RETURNING id_inventario;
                """
                inv_ins = db.execute_query(q_inv_ins, (id_sucursal, id_variante, cantidad, cantidad, stock_min), fetchone=True, commit=False)
                id_inventario = inv_ins[0]

            # H. Registrar Movimiento de Auditoría (t_movimiento_inventario)
            q_mov = f"""
            INSERT INTO {schema}.t_movimiento_inventario (
                id_inventario, id_usuario, tipo_movimiento, cantidad,
                stock_anterior, stock_nuevo, motivo, fecha_movimiento
            ) VALUES (%s, %s, 'ENTRADA_COMPRA', %s, %s, %s, %s, CURRENT_TIMESTAMP);
            """
            motivo_mov = f"Recepción de lote {numero_lote}" + (f" (OC #{id_orden_compra})" if id_orden_compra else "")
            db.execute_query(q_mov, (id_inventario, id_usuario, cantidad, stock_ant, stock_nuevo, motivo_mov), commit=False)

            items_procesados += 1

        # Commit final de toda la transacción
        db.conn.commit()

        # Registrar en Bitácora de Auditoría
        try:
            registrar_evento_db(
                id_usuario=id_usuario,
                usuario_nombre=token_data.get("nombre"),
                usuario_email=token_data.get("email") or token_data.get("sub"),
                id_empresa=id_empresa,
                id_sucursal=id_sucursal,
                modulo="INVENTARIO_COMPRAS",
                accion="IMPORTACION_MASIVA_LOTE",
                entidad="t_lote",
                id_entidad=str(id_lote),
                descripcion=f"Importación masiva exitosa del lote {numero_lote} ({total_prendas_lote} prendas) en sucursal {nombre_sucursal}.",
                resultado="EXITO",
                nivel="INFO",
                ip=None,
                user_agent=None,
                datos_anteriores=None,
                datos_nuevos={"total_prendas": total_prendas_lote, "numero_lote": numero_lote},
                metadatos={"id_lote": id_lote, "id_orden_compra": id_orden_compra},
                request_id=None
            )
        except Exception as bitacora_err:
            print(f"[AUDITORIA] Advertencia al registrar bitacora de importacion de lote: {bitacora_err}")

        return {
            "success": True,
            "mensaje": f"Lote {numero_lote} importado con éxito. Se registraron {total_prendas_lote} prendas en {nombre_sucursal}.",
            "id_lote": id_lote,
            "numero_lote": numero_lote,
            "id_orden_compra": id_orden_compra,
            "sucursal": nombre_sucursal,
            "total_prendas": total_prendas_lote,
            "costo_total": float(costo_total_lote),
            "productos_creados": productos_creados,
            "variantes_creadas": variantes_creadas,
            "items_procesados": items_procesados
        }
    except Exception as e:
        if db.conn:
            db.conn.rollback()
        raise e
    finally:
        db.close_connection()


# ==============================================================================
# 4. GESTIÓN DEL CICLO DE VIDA DE ÓRDENES DE COMPRA
# ==============================================================================

def listar_ordenes_compra_db(
    id_empresa: int,
    id_sucursal: Optional[int] = None,
    estado: Optional[str] = None
) -> List[Dict[str, Any]]:
    """Retorna las órdenes de compra de la empresa con filtros opcionales."""
    schema = _get_schema()
    db = PostgreSQL()
    db.create_connection()
    try:
        where = ["oc.id_empresa = %s"]
        params = [id_empresa]

        if id_sucursal:
            where.append("oc.id_sucursal = %s")
            params.append(id_sucursal)

        if estado and estado != "TODOS":
            where.append("UPPER(oc.estado) = %s")
            params.append(estado.upper())

        where_sql = " AND ".join(where)

        q = f"""
        SELECT 
            oc.id_orden_compra,
            oc.numero_orden,
            oc.id_sucursal,
            s.nombre AS sucursal,
            oc.id_proveedor,
            COALESCE(p.razon_social, 'Proveedor General') AS proveedor,
            oc.fecha_emision,
            oc.fecha_entrega_esperada,
            oc.estado,
            COALESCE(oc.total_estimado, 0) AS total_estimado,
            oc.observaciones,
            COALESCE(u_cr.nombre || ' ' || COALESCE(u_cr.apellido, ''), 'Usuario') AS creador,
            COALESCE(u_ap.nombre || ' ' || COALESCE(u_ap.apellido, ''), '') AS aprobador,
            oc.fecha_aprobacion,
            oc.motivo_rechazo,
            (SELECT COUNT(*) FROM {schema}.t_detalle_orden_compra doc WHERE doc.id_orden_compra = oc.id_orden_compra) AS total_items,
            (SELECT COALESCE(SUM(doc.cantidad_solicitada), 0) FROM {schema}.t_detalle_orden_compra doc WHERE doc.id_orden_compra = oc.id_orden_compra) AS total_prendas
        FROM {schema}.t_orden_compra oc
        JOIN {schema}.t_sucursal s ON s.id_sucursal = oc.id_sucursal
        LEFT JOIN {schema}.t_proveedor p ON p.id_proveedor = oc.id_proveedor
        LEFT JOIN {schema}.t_usuario u_cr ON u_cr.id_usuario = oc.id_usuario_creador
        LEFT JOIN {schema}.t_usuario u_ap ON u_ap.id_usuario = oc.id_usuario_aprobador
        WHERE {where_sql}
        ORDER BY oc.id_orden_compra DESC;
        """
        rows = db.execute_query(q, tuple(params), fetchall=True) or []
        res = []
        for r in rows:
            res.append({
                "id_orden_compra": r[0],
                "numero_orden": r[1],
                "id_sucursal": r[2],
                "sucursal": r[3],
                "id_proveedor": r[4],
                "proveedor": r[5],
                "fecha_emision": r[6].isoformat() if r[6] else None,
                "fecha_entrega_esperada": r[7].isoformat() if r[7] else None,
                "estado": r[8],
                "total_estimado": float(r[9]),
                "observaciones": r[10],
                "creador": r[11],
                "aprobador": r[12],
                "fecha_aprobacion": r[13].isoformat() if r[13] else None,
                "motivo_rechazo": r[14],
                "total_items": r[15],
                "total_prendas": r[16]
            })
        return res
    finally:
        db.close_connection()


def obtener_detalle_orden_compra_db(id_orden: int, id_empresa: int) -> Dict[str, Any]:
    """Obtiene la cabecera y el detalle de ítems de una orden de compra."""
    schema = _get_schema()
    db = PostgreSQL()
    db.create_connection()
    try:
        # Cabecera
        q_cab = f"""
        SELECT 
            oc.id_orden_compra, oc.numero_orden, oc.id_sucursal, s.nombre,
            oc.id_proveedor, COALESCE(p.razon_social, 'Sin Proveedor'),
            oc.fecha_emision, oc.fecha_entrega_esperada, oc.estado,
            oc.total_estimado, oc.observaciones, oc.motivo_rechazo
        FROM {schema}.t_orden_compra oc
        JOIN {schema}.t_sucursal s ON s.id_sucursal = oc.id_sucursal
        LEFT JOIN {schema}.t_proveedor p ON p.id_proveedor = oc.id_proveedor
        WHERE oc.id_orden_compra = %s AND oc.id_empresa = %s;
        """
        cab = db.execute_query(q_cab, (id_orden, id_empresa), fetchone=True)
        if not cab:
            raise ValueError(f"Orden de compra #{id_orden} no encontrada.")

        # Detalle de prendas
        q_det = f"""
        SELECT 
            doc.id_detalle_orden, doc.id_producto, doc.id_variante,
            doc.codigo_producto, doc.nombre_producto, doc.talla, doc.color,
            doc.sku, doc.cantidad_solicitada, doc.cantidad_recibida,
            doc.costo_unitario, doc.subtotal
        FROM {schema}.t_detalle_orden_compra doc
        WHERE doc.id_orden_compra = %s
        ORDER BY doc.id_detalle_orden ASC;
        """
        items_raw = db.execute_query(q_det, (id_orden,), fetchall=True) or []
        items = []
        for it in items_raw:
            items.append({
                "id_detalle_orden": it[0],
                "id_producto": it[1],
                "id_variante": it[2],
                "codigo_producto": it[3],
                "nombre_producto": it[4],
                "talla": it[5],
                "color": it[6],
                "sku": it[7],
                "cantidad_solicitada": it[8],
                "cantidad_recibida": it[9],
                "costo_unitario": float(it[10]),
                "subtotal": float(it[11])
            })

        return {
            "id_orden_compra": cab[0],
            "numero_orden": cab[1],
            "id_sucursal": cab[2],
            "sucursal": cab[3],
            "id_proveedor": cab[4],
            "proveedor": cab[5],
            "fecha_emision": cab[6].isoformat() if cab[6] else None,
            "fecha_entrega_esperada": cab[7].isoformat() if cab[7] else None,
            "estado": cab[8],
            "total_estimado": float(cab[9]),
            "observaciones": cab[10],
            "motivo_rechazo": cab[11],
            "items": items
        }
    finally:
        db.close_connection()


def aprobar_orden_compra_db(id_orden: int, token_data: Dict[str, Any]) -> Dict[str, Any]:
    """Aprueba una orden de compra pendiente para permitir su posterior recepción."""
    id_empresa = int(token_data.get("id_empresa") or 1)
    id_usuario = int(token_data.get("nro_usuario") or 1)
    schema = _get_schema()

    db = PostgreSQL()
    db.create_connection()
    try:
        # Resolver id_usuario seguro
        q_u = f"SELECT id_usuario FROM {schema}.t_usuario WHERE id_usuario = %s;"
        u_chk = db.execute_query(q_u, (id_usuario,), fetchone=True)
        if not u_chk:
            q_u_alt = f"SELECT id_usuario FROM {schema}.t_usuario WHERE id_empresa = %s LIMIT 1;"
            u_alt = db.execute_query(q_u_alt, (id_empresa,), fetchone=True)
            id_usuario = u_alt[0] if u_alt else None

        q_chk = f"SELECT id_orden_compra, estado, numero_orden FROM {schema}.t_orden_compra WHERE id_orden_compra = %s AND id_empresa = %s;"
        row = db.execute_query(q_chk, (id_orden, id_empresa), fetchone=True)
        if not row:
            raise ValueError(f"Orden de compra #{id_orden} no existe en la empresa.")
        if row[1] not in ["PENDIENTE_APROBACION", "BORRADOR"]:
            raise ValueError(f"No se puede aprobar la orden #{id_orden} porque ya se encuentra en estado '{row[1]}'.")

        q_upd = f"""
        UPDATE {schema}.t_orden_compra
        SET estado = 'APROBADA',
            id_usuario_aprobador = %s,
            fecha_aprobacion = CURRENT_TIMESTAMP
        WHERE id_orden_compra = %s;
        """
        db.execute_query(q_upd, (id_usuario, id_orden), commit=True)
        return {"success": True, "mensaje": f"Orden de compra {row[2]} aprobada exitosamente."}
    finally:
        db.close_connection()


def rechazar_orden_compra_db(id_orden: int, motivo: str, token_data: Dict[str, Any]) -> Dict[str, Any]:
    """Rechaza una orden de compra justificando el motivo."""
    id_empresa = int(token_data.get("id_empresa") or 1)
    schema = _get_schema()

    db = PostgreSQL()
    db.create_connection()
    try:
        q_chk = f"SELECT id_orden_compra, estado, numero_orden FROM {schema}.t_orden_compra WHERE id_orden_compra = %s AND id_empresa = %s;"
        row = db.execute_query(q_chk, (id_orden, id_empresa), fetchone=True)
        if not row:
            raise ValueError(f"Orden de compra #{id_orden} no existe en la empresa.")

        q_upd = f"""
        UPDATE {schema}.t_orden_compra
        SET estado = 'RECHAZADA',
            motivo_rechazo = %s
        WHERE id_orden_compra = %s;
        """
        db.execute_query(q_upd, (motivo, id_orden), commit=True)
        return {"success": True, "mensaje": f"Orden de compra {row[2]} rechazada."}
    finally:
        db.close_connection()


def recibir_mercaderia_orden_db(
    id_orden: int,
    recepcion_data: Dict[str, Any],
    token_data: Dict[str, Any]
) -> Dict[str, Any]:
    """
    Confirma la recepción física de mercadería de una orden aprobada:
    1. Registra el lote recibido en t_lote.
    2. Actualiza t_detalle_orden_compra con las cantidades recibidas.
    3. Incrementa stock en t_inventario para cada variante.
    4. Registra los movimientos de inventario en t_movimiento_inventario.
    5. Actualiza el estado de la orden a RECIBIDA_TOTAL.
    """
    id_empresa = int(token_data.get("id_empresa") or 1)
    id_usuario = int(token_data.get("nro_usuario") or 1)
    schema = _get_schema()

    db = PostgreSQL()
    db.create_connection()
    try:
        # Resolver id_usuario seguro
        q_u = f"SELECT id_usuario FROM {schema}.t_usuario WHERE id_usuario = %s;"
        u_chk = db.execute_query(q_u, (id_usuario,), fetchone=True)
        if not u_chk:
            q_u_alt = f"SELECT id_usuario FROM {schema}.t_usuario WHERE id_empresa = %s LIMIT 1;"
            u_alt = db.execute_query(q_u_alt, (id_empresa,), fetchone=True)
            id_usuario = u_alt[0] if u_alt else None

        # Validar orden
        q_oc = f"""
        SELECT id_orden_compra, id_sucursal, id_proveedor, estado, numero_orden, total_estimado
        FROM {schema}.t_orden_compra
        WHERE id_orden_compra = %s AND id_empresa = %s;
        """
        oc = db.execute_query(q_oc, (id_orden, id_empresa), fetchone=True)
        if not oc:
            raise ValueError(f"Orden #{id_orden} no encontrada.")
        if oc[3] != 'APROBADA':
            raise ValueError(f"Solo se pueden recepcionar órdenes en estado 'APROBADA'. Estado actual: '{oc[3]}'.")

        id_sucursal = oc[1]
        id_proveedor = oc[2]
        numero_oc = oc[4]

        # Obtener ítems de la orden
        q_items = f"""
        SELECT id_detalle_orden, id_producto, id_variante, cantidad_solicitada, costo_unitario, sku, nombre_producto
        FROM {schema}.t_detalle_orden_compra
        WHERE id_orden_compra = %s;
        """
        items = db.execute_query(q_items, (id_orden,), fetchall=True) or []
        if not items:
            raise ValueError("La orden de compra no tiene ítems registrados.")

        timestamp_lote = datetime.datetime.now().strftime("%Y%m%d_%H%M")
        numero_lote = recepcion_data.get("numero_lote") or f"LOT-{timestamp_lote}-{id_orden}"
        guia_remision = recepcion_data.get("guia_remision") or f"REC-{id_orden}"
        observaciones = recepcion_data.get("observaciones") or f"Recepción física de Orden {numero_oc}"

        total_prendas = sum(it[3] for it in items)
        costo_total = sum(Decimal(str(it[3])) * Decimal(str(it[4])) for it in items)

        # 1. Crear lote
        q_lote = f"""
        INSERT INTO {schema}.t_lote (
            id_empresa, id_sucursal, id_orden_compra, id_proveedor,
            numero_lote, guia_remision, total_prendas, costo_total_lote,
            id_usuario_receptor, observaciones
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id_lote;
        """
        l_res = db.execute_query(
            q_lote,
            (id_empresa, id_sucursal, id_orden, id_proveedor, numero_lote, guia_remision, total_prendas, float(costo_total), id_usuario, observaciones),
            fetchone=True, commit=False
        )
        id_lote = l_res[0]

        # 2. Afectar inventario por cada ítem
        for it in items:
            id_det = it[0]
            id_var = it[2]
            cant = it[3]

            # Actualizar cantidad recibida en detalle
            q_upd_det = f"UPDATE {schema}.t_detalle_orden_compra SET cantidad_recibida = %s WHERE id_detalle_orden = %s;"
            db.execute_query(q_upd_det, (cant, id_det), commit=False)

            if id_var:
                q_inv = f"SELECT id_inventario, stock_actual FROM {schema}.t_inventario WHERE id_sucursal = %s AND id_variante = %s LIMIT 1;"
                inv = db.execute_query(q_inv, (id_sucursal, id_var), fetchone=True)
                if inv:
                    id_inv = inv[0]
                    stk_ant = inv[1]
                    q_upd_inv = f"""
                    UPDATE {schema}.t_inventario
                    SET stock_actual = stock_actual + %s,
                        stock_disponible = stock_disponible + %s,
                        fecha_actualizacion = CURRENT_TIMESTAMP
                    WHERE id_inventario = %s;
                    """
                    db.execute_query(q_upd_inv, (cant, cant, id_inv), commit=False)
                else:
                    stk_ant = 0
                    q_ins_inv = f"""
                    INSERT INTO {schema}.t_inventario (
                        id_sucursal, id_variante, stock_actual, stock_reservado,
                        stock_disponible, stock_minimo, estado, fecha_actualizacion
                    ) VALUES (%s, %s, %s, 0, %s, 5, TRUE, CURRENT_TIMESTAMP)
                    RETURNING id_inventario;
                    """
                    inv_res = db.execute_query(q_ins_inv, (id_sucursal, id_var, cant, cant), fetchone=True, commit=False)
                    id_inv = inv_res[0]

                # Registrar movimiento de inventario
                q_mov = f"""
                INSERT INTO {schema}.t_movimiento_inventario (
                    id_inventario, id_usuario, tipo_movimiento, cantidad,
                    stock_anterior, stock_nuevo, motivo, fecha_movimiento
                ) VALUES (%s, %s, 'ENTRADA_COMPRA', %s, %s, %s, %s, CURRENT_TIMESTAMP);
                """
                db.execute_query(q_mov, (id_inv, id_usuario, cant, stk_ant, stk_ant + cant, f"Recepción de OC #{id_orden} - Lote {numero_lote}"), commit=False)

        # 3. Marcar orden como RECIBIDA_TOTAL
        q_upd_oc = f"UPDATE {schema}.t_orden_compra SET estado = 'RECIBIDA_TOTAL' WHERE id_orden_compra = %s;"
        db.execute_query(q_upd_oc, (id_orden,), commit=False)

        db.conn.commit()
        return {
            "success": True,
            "mensaje": f"Orden {numero_oc} recibida exitosamente. Lote #{id_lote} ({numero_lote}) registrado en inventario.",
            "id_lote": id_lote,
            "numero_lote": numero_lote,
            "total_prendas": total_prendas
        }
    except Exception as e:
        if db.conn:
            db.conn.rollback()
        raise e
    finally:
        db.close_connection()


def listar_lotes_db(id_empresa: int, id_sucursal: Optional[int] = None) -> List[Dict[str, Any]]:
    """Lista todos los lotes ingresados al inventario de la empresa."""
    schema = _get_schema()
    db = PostgreSQL()
    db.create_connection()
    try:
        where = ["l.id_empresa = %s"]
        params = [id_empresa]
        if id_sucursal:
            where.append("l.id_sucursal = %s")
            params.append(id_sucursal)

        where_sql = " AND ".join(where)

        q = f"""
        SELECT 
            l.id_lote,
            l.numero_lote,
            l.id_sucursal,
            s.nombre AS sucursal,
            l.id_orden_compra,
            COALESCE(oc.numero_orden, 'Recepción Directa') AS orden_compra,
            COALESCE(p.razon_social, 'Sin Proveedor') AS proveedor,
            l.fecha_ingreso,
            l.guia_remision,
            l.total_prendas,
            l.costo_total_lote,
            COALESCE(u.nombre || ' ' || COALESCE(u.apellido, ''), 'Usuario') AS receptor,
            l.observaciones
        FROM {schema}.t_lote l
        JOIN {schema}.t_sucursal s ON s.id_sucursal = l.id_sucursal
        LEFT JOIN {schema}.t_orden_compra oc ON oc.id_orden_compra = l.id_orden_compra
        LEFT JOIN {schema}.t_proveedor p ON p.id_proveedor = l.id_proveedor
        LEFT JOIN {schema}.t_usuario u ON u.id_usuario = l.id_usuario_receptor
        WHERE {where_sql}
        ORDER BY l.id_lote DESC;
        """
        rows = db.execute_query(q, tuple(params), fetchall=True) or []
        res = []
        for r in rows:
            res.append({
                "id_lote": r[0],
                "numero_lote": r[1],
                "id_sucursal": r[2],
                "sucursal": r[3],
                "id_orden_compra": r[4],
                "orden_compra": r[5],
                "proveedor": r[6],
                "fecha_ingreso": r[7].isoformat() if r[7] else None,
                "guia_remision": r[8],
                "total_prendas": r[9],
                "costo_total_lote": float(r[10] or 0),
                "receptor": r[11],
                "observaciones": r[12]
            })
        return res
    finally:
        db.close_connection()
